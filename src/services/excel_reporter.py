"""
Corporate Excel report generation service using OpenPyXL.
Applies executive KPI cards, B3 corporate palette, accounting number formats, and dynamic column autosizing.
"""
from pathlib import Path
from typing import Dict, Any
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class B3ExcelReporter:
    """Generates styled corporate Excel workbooks with KPI summaries and formatted data tables."""

    NAVY_B3 = "002060"
    GRAY_CARD = "F2F4F8"
    WHITE = "FFFFFF"

    def generate_report(
        self,
        df: pd.DataFrame,
        metrics: Dict[str, Any],
        output_path: str = "output/b3_ibovespa_carteira_teorica.xlsx",
    ) -> str:
        """Creates and styles an executive Excel workbook with KPI dashboard and asset composition."""
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Carteira Teórica IBOV"
        ws.views.sheetView[0].showGridLines = True

        # Top banner
        ws.merge_cells("A1:I1")
        banner = ws["A1"]
        banner.value = "B3 - CARTEIRA TEÓRICA DO ÍNDICE IBOVESPA"
        banner.font = Font(name="Calibri", size=14, bold=True, color=self.WHITE)
        banner.fill = PatternFill(start_color=self.NAVY_B3, end_color=self.NAVY_B3, fill_type="solid")
        banner.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # KPI Cards
        cards_data = [
            ("B3", "B4", "TOTAL ATIVOS", f"{metrics.get('total_assets', 0)}"),
            ("D3", "D4", "QUANTIDADE DE AÇÕES", f"{metrics.get('total_quantity', 0):,}".replace(",", ".")),
            ("F3", "F4", "TOP 5 CONCENTRAÇÃO", f"{metrics.get('top_5_share_pct', 0.0):.2f}%"),
            ("H3", "H4", "TOP 10 CONCENTRAÇÃO", f"{metrics.get('top_10_share_pct', 0.0):.2f}%"),
        ]

        card_title_font = Font(name="Calibri", size=9, bold=True, color="595959")
        card_value_font = Font(name="Calibri", size=13, bold=True, color=self.NAVY_B3)
        card_fill = PatternFill(start_color=self.GRAY_CARD, end_color=self.GRAY_CARD, fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")

        for pos_title, pos_val, title, val in cards_data:
            c_title, c_val = ws[pos_title], ws[pos_val]
            c_title.value, c_title.font, c_title.fill, c_title.alignment = title, card_title_font, card_fill, center_align
            c_val.value, c_val.font, c_val.fill, c_val.alignment = val, card_value_font, card_fill, center_align

        # Data table header
        headers = list(df.columns)
        header_font = Font(name="Calibri", size=11, bold=True, color=self.WHITE)
        header_fill = PatternFill(start_color=self.NAVY_B3, end_color=self.NAVY_B3, fill_type="solid")

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=6, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
        ws.row_dimensions[6].height = 25

        # Insertion and formatting of data rows
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )

        for row_idx, row_data in enumerate(df.values, start=7):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                col_name = headers[col_idx - 1]

                # Applying accounting formats and standardized alignments
                if col_name == "Qtde. Teórica":
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_name in ("Part. (%)", "Part. Acumulada (%)"):
                    cell.number_format = '0.000"%"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_name in ("Posição", "Código", "Tipo", "Classe Ação", "Classificação"):
                    cell.alignment = center_align
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Column width autofit (evaluated only from row 6 downwards to avoid banner distortion)
        for col_idx, header in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                len(str(ws.cell(row=r, column=col_idx).value or ""))
                for r in range(6, len(df) + 7)
            )
            ws.column_dimensions[col_letter].width = max(max_len + 4, 10)

        wb.save(output_path)
        print(f"[SUCCESS] [B3ExcelReporter] Corporate Excel report generated: {output_path}")
        return output_path
